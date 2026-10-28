"""Inventory report writers: flattening rows and the Excel workbook."""

from __future__ import annotations

from collections import Counter
from collections.abc import Sequence
from pathlib import Path
from typing import Any

from ..errors import LzctlError
from . import BASE_COLUMNS, Collector, Resource


def flatten(resource: Resource, tag_columns: Sequence[str]) -> dict[str, Any]:
    """One flat row: tag columns become ``tag:<Key>``, remaining tags a k=v string."""
    tags = resource.get("tags") or {}
    row = {k: v for k, v in resource.items() if k != "tags"}
    for key in tag_columns:
        row[f"tag:{key}"] = tags.get(key)
    row["tags"] = ";".join(f"{k}={v}" for k, v in sorted(tags.items()) if k not in tag_columns) or None
    return row


def columns_for(collector: Collector | None, tag_columns: Sequence[str]) -> list[str]:
    extra = list(collector.columns) if collector else []
    return [*BASE_COLUMNS, *extra, *(f"tag:{k}" for k in tag_columns), "tags"]


def write_xlsx(
    path: Path,
    resources: Sequence[Resource],
    collectors: Sequence[Collector],
    tag_columns: Sequence[str],
    title: str,
) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as err:
        raise LzctlError("Excel output needs openpyxl: pip install 'lzctl[excel]'") from err

    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDE3EA")
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    summary.append([title])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([])
    counts = Counter((r["account"], r["service"]) for r in resources)
    accounts = sorted({r["account"] for r in resources})
    services = [c.name for c in collectors]
    summary.append(["account", *services, "total"])
    for cell in summary[3]:
        cell.font, cell.fill = bold, header_fill
    for acct in accounts:
        values = [counts.get((acct, s), 0) for s in services]
        summary.append([acct, *values, sum(values)])
    totals = [sum(counts.get((a, s), 0) for a in accounts) for s in services]
    summary.append(["total", *totals, sum(totals)])
    for cell in summary[summary.max_row]:
        cell.font = bold

    def fill_sheet(ws, rows, columns):
        ws.append(columns)
        for cell in ws[1]:
            cell.font, cell.fill = bold, header_fill
        for row in rows:
            ws.append([_excel_value(row.get(c)) for c in columns])
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(columns, 1):
            width = max([len(str(col))] + [len(str(r.get(col) or "")) for r in rows[:500]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)

    for coll in collectors:
        rows = [flatten(r, tag_columns) for r in resources if r["service"] == coll.name]
        fill_sheet(wb.create_sheet(coll.name[:31]), rows, columns_for(coll, tag_columns))

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _excel_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)
